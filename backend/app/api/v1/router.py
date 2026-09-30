import csv
import io
import json
from datetime import UTC, date, datetime, timedelta

from fastapi import APIRouter, Depends, File, HTTPException, Query, UploadFile
from fastapi.responses import FileResponse, StreamingResponse
from openpyxl import Workbook, load_workbook
from sqlalchemy import func, or_, select
from sqlalchemy.exc import IntegrityError
from sqlalchemy.orm import Session

from app.core.security import create_access_token, hash_password, verify_password
from app.db.session import get_db
from app.models import (
    AuditLog, AuthSession, ComplianceAssignment, ComplianceComment, ComplianceEvidence, ComplianceImport, InAppNotification,
    ComplianceImportError, ComplianceImportRow, ComplianceInstance, ComplianceRule, ComplianceRuleVersion, ComplianceSubmission,
    ComplianceWorkflowHistory, Contractor, ContractorSite, IndustryType, Organization, RuleApplicabilityCriteria,
    State, Unit, User, UserRoleScope,
)
from app.schemas import (
    ActivityInput, AssignmentInput, CommentInput, ContractorInput, ContractorSiteInput, GenerationRequest,
    LoginRequest, OrganizationInput, PlatformOrganizationInput, OrganizationStatusInput, PasswordChangeRequest, RoleScopeInput, RuleInput, UnitInput, UserInput,
    WorkflowDecision, UserUpdateInput, EvidenceVerificationInput, NotApplicableInput, RuleStatusInput,
)
from app.services.access import CHECKER_ROLES, MAKER_ROLES, MANAGER_ROLES, PLATFORM_ADMIN_ROLE, can_access_subject, can_read_audit, current_user, is_org_admin, require_org_admin, require_org_admin_for_entity_management, roles_for, scopes_for, user_summary
from app.services.audit import audit
from app.services.compliance import SUBJECTS, assert_assignment_or_admin, generate_for_subject, submit, transition
from app.services.serialization import entity_dict, instance_dict, model_dict
from app.storage.files import delete_stored_file, file_path, save_upload

router = APIRouter()


def org_for(user: User) -> str:
    organization_id = getattr(user, "_morax_active_organization_id", user.organization_id)
    if not organization_id:
        raise HTTPException(409, "User is not assigned to an organization")
    return organization_id


def require_platform_admin(user: User) -> None:
    if user.platform_role != PLATFORM_ADMIN_ROLE:
        raise HTTPException(403, "MORAX platform administrator permission is required")


def organization_dict(organization: Organization) -> dict:
    fields = [field.name for field in organization.__table__.columns]
    result = model_dict(organization, fields)
    result["organization_name"] = result.pop("name")
    result["organization_code"] = result.pop("code")
    result["registered_address"] = result.pop("address")
    result["primary_contact_email"] = result.pop("contact_email")
    result["primary_contact_phone"] = result.pop("contact_phone")
    return result


def apply_platform_organization(organization: Organization, payload: PlatformOrganizationInput, actor_id: str) -> None:
    values = payload.model_dump()
    organization.name = values["organization_name"].strip()
    organization.code = values["organization_code"].strip().upper()
    organization.legal_name = values["legal_name"]
    organization.registration_number = values["registration_number"]
    organization.pan = values["pan"]
    organization.gstin = values["gstin"]
    organization.address = values["registered_address"]
    organization.city = values["city"]
    organization.state_id = values["state_id"]
    organization.pincode = values["pincode"]
    organization.primary_contact_name = values["primary_contact_name"]
    organization.contact_email = str(values["primary_contact_email"]) if values["primary_contact_email"] else None
    organization.contact_phone = values["primary_contact_phone"]
    organization.compliance_start_date = values["compliance_start_date"]
    organization.status = values["status"].upper()
    organization.updated_by_id = actor_id


def require_manager(db: Session, user: User, organization_id: str) -> None:
    if not is_org_admin(db, user, organization_id) and not (roles_for(db, user.id) & MANAGER_ROLES):
        raise HTTPException(403, "Administrative permission is required")


def require_deletable_entity(db: Session, subject_type: str, subject_id: str) -> None:
    """Never sever historical compliance records from the entity they describe."""
    foreign_key = {
        "UNIT": ComplianceInstance.unit_id,
        "CONTRACTOR": ComplianceInstance.contractor_id,
        "CONTRACTOR_SITE": ComplianceInstance.contractor_site_id,
    }.get(subject_type)
    if foreign_key is None:
        raise HTTPException(422, "Unsupported entity type")
    history = db.scalar(select(ComplianceInstance.id).where(
        (foreign_key == subject_id)
        | ((ComplianceInstance.subject_type == subject_type) & (ComplianceInstance.subject_id == subject_id))
    ).limit(1))
    if history:
        label = {"UNIT": "Unit", "CONTRACTOR": "Contractor", "CONTRACTOR_SITE": "Contractor site"}[subject_type]
        raise HTTPException(409, f"Cannot delete {label} because it has compliance history. Set its status to INACTIVE instead.")


def commit_or_conflict(db: Session, detail: str) -> None:
    """Turn expected unique-key races into a usable API validation response."""
    try:
        db.commit()
    except IntegrityError as exc:
        db.rollback()
        raise HTTPException(409, detail) from exc


def require_scope_belongs_to_organization(
    db: Session, organization_id: str, scope_type: str, scope_id: str
) -> None:
    if scope_type == "ORGANIZATION":
        if scope_id != organization_id:
            raise HTTPException(422, "Organization scope must be the active organization")
        return
    model = {
        "UNIT": Unit,
        "CONTRACTOR": Contractor,
        "CONTRACTOR_SITE": ContractorSite,
    }.get(scope_type)
    if not model:
        raise HTTPException(422, "Unsupported role scope type")
    subject = db.get(model, scope_id)
    if not subject or subject.organization_id != organization_id:
        raise HTTPException(422, "Role scope must belong to the active organization")


def validate_industry_selection(db: Session, industry_type_id: str | None, other_industry_name: str | None) -> str | None:
    """Validate the controlled industry category and its optional custom label.

    A free-text label is permitted only under the seeded OTHER category. This
    keeps rule applicability governed by Compliance Master rather than allowing
    arbitrary industries to silently change compliance behaviour.
    """
    custom_name = other_industry_name.strip() if other_industry_name else None
    if custom_name == "":
        custom_name = None
    industry = db.get(IndustryType, industry_type_id) if industry_type_id else None
    if industry_type_id and not industry:
        raise HTTPException(422, "Industry type must be a configured master value")
    is_other = bool(industry and industry.code.upper() == "OTHER")
    if is_other and not custom_name:
        raise HTTPException(422, "Enter the industry name when Other industry is selected")
    if not is_other and custom_name:
        raise HTTPException(422, "A custom industry name is allowed only for Other industry")
    return custom_name


def list_page(items: list, page: int, page_size: int) -> dict:
    total = len(items)
    offset = (page - 1) * page_size
    return {"items": items[offset:offset + page_size], "total": total, "page": page, "page_size": page_size}


def entity_page(
    db: Session,
    model,
    entity_type: str,
    organization_id: str,
    page: int,
    page_size: int,
    q: str | None = None,
    status: str | None = None,
    sort_by: str = "name",
    sort_dir: str = "asc",
) -> dict:
    """Return a bounded, tenant-scoped page for the entity management grids.

    The public entity endpoints still support their legacy list response where
    required. New UI consumers opt into this paged path, keeping the API
    compatible while avoiding unbounded record transfers as organizations grow.
    """
    filters = [model.organization_id == organization_id]
    if q:
        text = f"%{q.strip()}%"
        filters.append(or_(model.name.ilike(text), model.code.ilike(text)))
    if status:
        filters.append(model.status == status.upper())
    allowed_sorts = {"name", "code", "status", "city"}
    column = getattr(model, sort_by if sort_by in allowed_sorts else "name")
    ordering = column.desc() if sort_dir.lower() == "desc" else column.asc()
    safe_page = max(page, 1)
    safe_size = min(max(page_size, 1), 100)
    total = db.scalar(select(func.count()).select_from(model).where(*filters)) or 0
    rows = db.scalars(
        select(model)
        .where(*filters)
        .order_by(ordering, model.id)
        .offset((safe_page - 1) * safe_size)
        .limit(safe_size)
    ).all()
    return {
        "items": [entity_dict(db, row, entity_type) for row in rows],
        "total": total,
        "page": safe_page,
        "page_size": safe_size,
    }


def parse_bool(value, default=True):
    if value is None or value == "":
        return default
    return str(value).strip().lower() in {"true", "1", "yes", "y", "active"}


def stage_construction_checklist(workbook, organization_id: str, user: User, db: Session) -> dict:
    """Stage COW sample rows as Master records; no statutory schedule is inferred."""
    sheet = workbook["COW"]
    headers = [cell.value for cell in next(sheet.iter_rows(min_row=1, max_row=1))]
    required = {"Sl. No.", "Code / Act & Rules", "Procedural Compliance Check", "Frequency"}
    if not required.issubset(set(headers)):
        raise HTTPException(422, "COW worksheet does not contain the expected construction checklist headers")
    record = ComplianceImport(organization_id=organization_id, uploaded_by_id=user.id, original_filename="construction-checklist.xlsx")
    db.add(record); db.flush()
    frequency_map = {"daily": "DAILY", "weekly": "WEEKLY", "monthly": "MONTHLY", "quarterly": "QUARTERLY", "half-yearly": "HALF_YEARLY", "annually": "ANNUAL", "annual": "ANNUAL", "one-time": "ONE_TIME", "event-based": "EVENT_BASED", "periodic": "EVENT_BASED", "ongoing": "EVENT_BASED", "continuous": "EVENT_BASED", "joining / change": "EVENT_BASED"}
    total = valid = errors = 0
    for row_number, values in enumerate(sheet.iter_rows(min_row=2, values_only=True), start=2):
        if not any(value is not None and str(value).strip() for value in values):
            continue
        total += 1
        row = {headers[index]: values[index] for index in range(len(headers))}
        messages = []
        if not row.get("Procedural Compliance Check"): messages.append(("Procedural Compliance Check", "Value is required"))
        frequency = frequency_map.get(str(row.get("Frequency") or "").strip().lower())
        if not frequency: messages.append(("Frequency", "Unsupported frequency"))
        payload = {
            "Compliance ID": f"COW-{int(row['Sl. No.']):03d}" if isinstance(row.get("Sl. No."), (int, float)) else f"COW-{row_number:03d}",
            "Compliance Name": row.get("Procedural Compliance Check") or "", "Version": 1, "Entity Type": "UNIT",
            "Frequency": frequency or "EVENT_BASED", "Due Date Rule": "MANUAL", "Due Date Offset": None,
            "Effective From": date.today().isoformat(), "Active": "TRUE", "Act": row.get("Code / Act & Rules"),
            "Section": row.get("Section / Rule Provision"), "Legal Description": row.get("What the Act / Rule Says"),
            "Form Number": row.get("Form / Notice / Application"), "Risk Level": row.get("Risk Category") or "MEDIUM",
            "Required Document": row.get("Evidence to Verify"), "Description": "Imported from the construction industry COW checklist. Configure a due-date rule before automatic generation.",
        }
        item = ComplianceImportRow(import_id=record.id, row_number=row_number, payload_json=json.dumps(payload, default=str), valid=not messages)
        db.add(item); db.flush()
        for field, message in messages: db.add(ComplianceImportError(import_row_id=item.id, field=field, message=message))
        if messages: errors += 1
        else: valid += 1
    record.total_rows = total; record.valid_rows = valid; record.error_rows = errors
    audit(db, organization_id=organization_id, actor_id=user.id, action="VALIDATE_CONSTRUCTION_IMPORT", module="COMPLIANCE_MASTER", entity_type="ComplianceImport", entity_id=record.id, new={"total": total, "valid": valid, "errors": errors})
    db.commit()
    return {"id": record.id, "total_rows": total, "valid_rows": valid, "error_rows": errors}


@router.post("/auth/login")
def login(payload: LoginRequest, db: Session = Depends(get_db)):
    user = db.scalar(select(User).where(User.email == payload.email.lower()))
    if not user or not user.active or not verify_password(payload.password, user.password_hash):
        raise HTTPException(401, "Invalid email or password")
    organization = db.get(Organization, user.organization_id) if user.organization_id else None
    if user.platform_role != PLATFORM_ADMIN_ROLE and (not organization or organization.status == "INACTIVE"):
        raise HTTPException(403, "Your organization is inactive")
    session = AuthSession(user_id=user.id, active_organization_id=user.organization_id, expires_at=datetime.now(UTC) + timedelta(minutes=480))
    db.add(session)
    audit(db, organization_id=user.organization_id, actor_id=user.id, action="LOGIN", module="AUTH", entity_type="User", entity_id=user.id)
    db.commit()
    return {"access_token": create_access_token(user.id, session.id), "token_type": "bearer", "user": user_summary(db, user)}


@router.post("/auth/impersonate/{user_id}")
def start_impersonation(user_id: str, user: User = Depends(current_user), db: Session = Depends(get_db)):
    """Create a temporary, server-linked session to validate a tenant user's access."""
    organization_id = org_for(user)
    require_org_admin(db, user, organization_id)
    if getattr(user, "_morax_impersonated_by_id", None):
        raise HTTPException(409, "End the current impersonation before starting another one")
    target = db.get(User, user_id)
    if not target or target.organization_id != organization_id:
        raise HTTPException(404, "User not found")
    if not target.active:
        raise HTTPException(422, "Inactive users cannot be impersonated")
    if target.platform_role:
        raise HTTPException(403, "Platform accounts cannot be impersonated from an organization")
    if target.id == user.id:
        raise HTTPException(422, "You are already signed in as this user")

    originating_session_id = getattr(user, "_morax_session_id", None)
    originating_session = db.get(AuthSession, originating_session_id)
    if not originating_session or originating_session.user_id != user.id:
        raise HTTPException(401, "Your administrator session is no longer available")
    session = AuthSession(
        user_id=target.id,
        active_organization_id=organization_id,
        impersonated_by_session_id=originating_session.id,
        impersonated_by_id=user.id,
        expires_at=datetime.now(UTC) + timedelta(minutes=30),
    )
    db.add(session)
    audit(
        db,
        organization_id=organization_id,
        actor_id=user.id,
        action="IMPERSONATION_STARTED",
        module="AUTH",
        entity_type="User",
        entity_id=target.id,
        new={"target_email": target.email, "target_roles": sorted(roles_for(db, target.id))},
    )
    db.commit()
    setattr(target, "_morax_active_organization_id", organization_id)
    setattr(target, "_morax_session_id", session.id)
    setattr(target, "_morax_impersonated_by_id", user.id)
    setattr(target, "_morax_impersonated_by_session_id", originating_session.id)
    return {
        "access_token": create_access_token(target.id, session.id),
        "token_type": "bearer",
        "user": user_summary(db, target),
    }


@router.post("/auth/end-impersonation")
def end_impersonation(user: User = Depends(current_user), db: Session = Depends(get_db)):
    impersonated_session = db.get(AuthSession, getattr(user, "_morax_session_id", None))
    originating_session_id = getattr(user, "_morax_impersonated_by_session_id", None)
    originating_user_id = getattr(user, "_morax_impersonated_by_id", None)
    if not impersonated_session or not originating_session_id or not originating_user_id:
        raise HTTPException(409, "There is no active impersonation to end")
    originating_session = db.get(AuthSession, originating_session_id)
    originating_user = db.get(User, originating_user_id)
    if (
        not originating_session
        or originating_session.revoked_at
        or originating_session.user_id != originating_user_id
        or not originating_user
        or not originating_user.active
    ):
        raise HTTPException(401, "The originating administrator session is no longer available")

    impersonated_session.revoked_at = datetime.now(UTC)
    audit(
        db,
        organization_id=org_for(user),
        actor_id=originating_user.id,
        action="IMPERSONATION_ENDED",
        module="AUTH",
        entity_type="User",
        entity_id=user.id,
        new={"target_email": user.email},
    )
    db.commit()
    setattr(originating_user, "_morax_active_organization_id", originating_session.active_organization_id)
    setattr(originating_user, "_morax_session_id", originating_session.id)
    return {
        "access_token": create_access_token(originating_user.id, originating_session.id),
        "token_type": "bearer",
        "user": user_summary(db, originating_user),
    }


@router.post("/auth/logout")
def logout(user: User = Depends(current_user), db: Session = Depends(get_db)):
    # Tokens are session-backed; revoke all current user sessions on explicit logout for predictable security.
    for session in db.scalars(select(AuthSession).where(AuthSession.user_id == user.id, AuthSession.revoked_at.is_(None))):
        session.revoked_at = datetime.now(UTC)
    audit(db, organization_id=user.organization_id, actor_id=user.id, action="LOGOUT", module="AUTH", entity_type="User", entity_id=user.id)
    db.commit()
    return {"ok": True}


@router.get("/auth/me")
def me(user: User = Depends(current_user), db: Session = Depends(get_db)):
    return user_summary(db, user)


@router.get("/platform/organizations")
def list_platform_organizations(q: str | None = None, status: str | None = None, user: User = Depends(current_user), db: Session = Depends(get_db)):
    require_platform_admin(user)
    query = select(Organization)
    if q:
        query = query.where(or_(Organization.name.ilike(f"%{q}%"), Organization.code.ilike(f"%{q}%"), Organization.legal_name.ilike(f"%{q}%")))
    if status:
        query = query.where(Organization.status == status.upper())
    return [organization_dict(item) for item in db.scalars(query.order_by(Organization.name))]


@router.post("/platform/organizations")
def create_platform_organization(payload: PlatformOrganizationInput, user: User = Depends(current_user), db: Session = Depends(get_db)):
    require_platform_admin(user)
    code = payload.organization_code.strip().upper()
    if db.scalar(select(Organization).where(Organization.code == code)):
        raise HTTPException(409, "Organization code already exists")
    if payload.state_id and not db.get(State, payload.state_id):
        raise HTTPException(422, "State must be a configured master value")
    organization = Organization(name=payload.organization_name.strip(), code=code, created_by_id=user.id)
    apply_platform_organization(organization, payload, user.id)
    db.add(organization); db.flush()
    audit(db, organization_id=organization.id, actor_id=user.id, action="ORGANIZATION_CREATED", module="PLATFORM", entity_type="Organization", entity_id=organization.id, new={"code": organization.code, "name": organization.name, "status": organization.status})
    db.commit()
    return organization_dict(organization)


@router.get("/platform/organizations/{organization_id}")
def get_platform_organization(organization_id: str, user: User = Depends(current_user), db: Session = Depends(get_db)):
    require_platform_admin(user)
    organization = db.get(Organization, organization_id)
    if not organization:
        raise HTTPException(404, "Organization not found")
    return organization_dict(organization)


@router.patch("/platform/organizations/{organization_id}")
def update_platform_organization(organization_id: str, payload: PlatformOrganizationInput, user: User = Depends(current_user), db: Session = Depends(get_db)):
    require_platform_admin(user)
    organization = db.get(Organization, organization_id)
    if not organization:
        raise HTTPException(404, "Organization not found")
    code = payload.organization_code.strip().upper()
    other = db.scalar(select(Organization).where(Organization.code == code, Organization.id != organization.id))
    if other:
        raise HTTPException(409, "Organization code already exists")
    if payload.state_id and not db.get(State, payload.state_id):
        raise HTTPException(422, "State must be a configured master value")
    old = {"code": organization.code, "name": organization.name, "status": organization.status}
    apply_platform_organization(organization, payload, user.id)
    audit(db, organization_id=organization.id, actor_id=user.id, action="ORGANIZATION_UPDATED", module="PLATFORM", entity_type="Organization", entity_id=organization.id, old=old, new={"code": organization.code, "name": organization.name, "status": organization.status})
    db.commit()
    return organization_dict(organization)


@router.post("/platform/organizations/{organization_id}/status")
def set_platform_organization_status(organization_id: str, payload: OrganizationStatusInput, user: User = Depends(current_user), db: Session = Depends(get_db)):
    require_platform_admin(user)
    organization = db.get(Organization, organization_id)
    if not organization:
        raise HTTPException(404, "Organization not found")
    old = organization.status; organization.status = payload.status
    organization.updated_by_id = user.id
    action = "ORGANIZATION_ACTIVATED" if payload.status == "ACTIVE" else "ORGANIZATION_DEACTIVATED"
    audit(db, organization_id=organization.id, actor_id=user.id, action=action, module="PLATFORM", entity_type="Organization", entity_id=organization.id, old={"status": old}, new={"status": organization.status})
    db.commit()
    return organization_dict(organization)


@router.post("/platform/organizations/{organization_id}/enter")
def enter_platform_organization(organization_id: str, user: User = Depends(current_user), db: Session = Depends(get_db)):
    require_platform_admin(user)
    organization = db.get(Organization, organization_id)
    if not organization:
        raise HTTPException(404, "Organization not found")
    session = db.get(AuthSession, getattr(user, "_morax_session_id", None))
    if not session:
        raise HTTPException(401, "Session is no longer active")
    session.active_organization_id = organization.id
    setattr(user, "_morax_active_organization_id", organization.id)
    audit(db, organization_id=organization.id, actor_id=user.id, action="ENTER_ORGANIZATION", module="PLATFORM", entity_type="Organization", entity_id=organization.id, new={"active_organization_id": organization.id})
    db.commit()
    return {"organization": organization_dict(organization), "user": user_summary(db, user)}


@router.post("/auth/change-password")
def change_password(payload: PasswordChangeRequest, user: User = Depends(current_user), db: Session = Depends(get_db)):
    if not verify_password(payload.current_password, user.password_hash):
        raise HTTPException(422, "Current password is incorrect")
    user.password_hash = hash_password(payload.new_password)
    user.must_change_password = False
    audit(db, organization_id=user.organization_id, actor_id=user.id, action="CHANGE_PASSWORD", module="AUTH", entity_type="User", entity_id=user.id)
    db.commit()
    return {"ok": True}


@router.get("/master-data/states")
def states(user: User = Depends(current_user), db: Session = Depends(get_db)):
    return [model_dict(row, ["id", "name", "code", "active"]) for row in db.scalars(select(State).where(State.active.is_(True)).order_by(State.name))]


@router.get("/master-data/industry-types")
def industries(user: User = Depends(current_user), db: Session = Depends(get_db)):
    return [model_dict(row, ["id", "name", "code", "active"]) for row in db.scalars(select(IndustryType).where(IndustryType.active.is_(True)).order_by(IndustryType.name))]


@router.get("/organizations")
def get_organization(user: User = Depends(current_user), db: Session = Depends(get_db)):
    organization = db.get(Organization, org_for(user))
    if not organization:
        raise HTTPException(404, "Organization not found")
    return model_dict(organization, [field.name for field in organization.__table__.columns])


@router.patch("/organizations")
def update_organization(payload: OrganizationInput, user: User = Depends(current_user), db: Session = Depends(get_db)):
    organization_id = org_for(user)
    require_org_admin(db, user, organization_id)
    organization = db.get(Organization, organization_id)
    old = {"name": organization.name, "status": organization.status}
    for key, value in payload.model_dump().items():
        setattr(organization, key, value)
    audit(db, organization_id=organization_id, actor_id=user.id, action="UPDATE_ORGANIZATION", module="ENTITIES", entity_type="Organization", entity_id=organization.id, old=old, new={"name": organization.name, "status": organization.status})
    db.commit()
    return model_dict(organization, [field.name for field in organization.__table__.columns])


@router.get("/units")
def list_units(page: int = 1, page_size: int = 50, q: str | None = None, status: str | None = None, sort_by: str = "name", sort_dir: str = "asc", user: User = Depends(current_user), db: Session = Depends(get_db)):
    organization_id = org_for(user)
    require_org_admin_for_entity_management(db, user, organization_id)
    return entity_page(db, Unit, "UNIT", organization_id, page, page_size, q, status, sort_by, sort_dir)


@router.get("/units/{unit_id}")
def get_unit(unit_id: str, user: User = Depends(current_user), db: Session = Depends(get_db)):
    unit = db.get(Unit, unit_id)
    if not unit or unit.organization_id != org_for(user):
        raise HTTPException(404, "Unit not found")
    require_org_admin_for_entity_management(db, user, unit.organization_id)
    return entity_dict(db, unit, "UNIT")


@router.post("/units")
def create_unit(payload: UnitInput, user: User = Depends(current_user), db: Session = Depends(get_db)):
    organization_id = org_for(user)
    require_org_admin_for_entity_management(db, user, organization_id)
    if db.scalar(select(Unit).where(Unit.organization_id == organization_id, Unit.code == payload.code)):
        raise HTTPException(409, "A unit with this code already exists in the organization")
    if not db.get(State, payload.state_id) or not db.get(IndustryType, payload.industry_type_id):
        raise HTTPException(422, "State and industry type must be configured master values")
    payload.other_industry_name = validate_industry_selection(db, payload.industry_type_id, payload.other_industry_name)
    unit = Unit(organization_id=organization_id, **payload.model_dump())
    db.add(unit)
    db.flush()
    generate_for_subject(db, organization_id, "UNIT", unit.id, date.today(), user.id)
    audit(db, organization_id=organization_id, actor_id=user.id, action="CREATE_UNIT", module="ENTITIES", entity_type="Unit", entity_id=unit.id, new={"name": unit.name, "code": unit.code})
    commit_or_conflict(db, "A unit with this code already exists in the organization")
    return entity_dict(db, unit, "UNIT")


@router.patch("/units/{unit_id}")
def update_unit(unit_id: str, payload: UnitInput, user: User = Depends(current_user), db: Session = Depends(get_db)):
    unit = db.get(Unit, unit_id)
    if not unit or unit.organization_id != org_for(user):
        raise HTTPException(404, "Unit not found")
    require_org_admin_for_entity_management(db, user, unit.organization_id)
    if db.scalar(select(Unit).where(Unit.organization_id == unit.organization_id, Unit.code == payload.code, Unit.id != unit.id)):
        raise HTTPException(409, "A unit with this code already exists in the organization")
    if not db.get(State, payload.state_id):
        raise HTTPException(422, "State must be a configured master value")
    payload.other_industry_name = validate_industry_selection(db, payload.industry_type_id, payload.other_industry_name)
    old = {"industry_type_id": unit.industry_type_id, "other_industry_name": unit.other_industry_name, "state_id": unit.state_id, "status": unit.status}
    for key, value in payload.model_dump().items(): setattr(unit, key, value)
    generate_for_subject(db, unit.organization_id, "UNIT", unit.id, date.today(), user.id)
    audit(db, organization_id=unit.organization_id, actor_id=user.id, action="UPDATE_UNIT", module="ENTITIES", entity_type="Unit", entity_id=unit.id, old=old, new={"industry_type_id": unit.industry_type_id, "other_industry_name": unit.other_industry_name, "state_id": unit.state_id, "status": unit.status})
    commit_or_conflict(db, "A unit with this code already exists in the organization")
    return entity_dict(db, unit, "UNIT")


@router.delete("/units/{unit_id}")
def delete_unit(unit_id: str, user: User = Depends(current_user), db: Session = Depends(get_db)):
    unit = db.get(Unit, unit_id)
    if not unit or unit.organization_id != org_for(user):
        raise HTTPException(404, "Unit not found")
    require_org_admin_for_entity_management(db, user, unit.organization_id)
    require_deletable_entity(db, "UNIT", unit.id)
    for contractor in db.scalars(select(Contractor).where(Contractor.unit_id == unit.id)):
        contractor.unit_id = None
    for site in db.scalars(select(ContractorSite).where(ContractorSite.unit_id == unit.id)):
        site.unit_id = None
    audit(db, organization_id=unit.organization_id, actor_id=user.id, action="DELETE_UNIT", module="ENTITIES", entity_type="Unit", entity_id=unit.id, old={"name": unit.name, "code": unit.code})
    db.delete(unit)
    db.commit()
    return {"ok": True}


@router.get("/contractors")
def list_contractors(paginated: bool = False, page: int = 1, page_size: int = 50, q: str | None = None, status: str | None = None, sort_by: str = "name", sort_dir: str = "asc", user: User = Depends(current_user), db: Session = Depends(get_db)):
    org = org_for(user)
    require_org_admin_for_entity_management(db, user, org)
    if paginated:
        return entity_page(db, Contractor, "CONTRACTOR", org, page, page_size, q, status, sort_by, sort_dir)
    return [entity_dict(db, row, "CONTRACTOR") for row in db.scalars(select(Contractor).where(Contractor.organization_id == org).order_by(Contractor.name))]


@router.get("/contractors/{contractor_id}")
def get_contractor(contractor_id: str, user: User = Depends(current_user), db: Session = Depends(get_db)):
    contractor = db.get(Contractor, contractor_id)
    if not contractor or contractor.organization_id != org_for(user):
        raise HTTPException(404, "Contractor not found")
    require_org_admin_for_entity_management(db, user, contractor.organization_id)
    return entity_dict(db, contractor, "CONTRACTOR")


@router.post("/contractors")
def create_contractor(payload: ContractorInput, user: User = Depends(current_user), db: Session = Depends(get_db)):
    org = org_for(user); require_org_admin_for_entity_management(db, user, org)
    if db.scalar(select(Contractor).where(Contractor.organization_id == org, Contractor.code == payload.code)):
        raise HTTPException(409, "A contractor with this code already exists in the organization")
    if payload.unit_id and (not db.get(Unit, payload.unit_id) or db.get(Unit, payload.unit_id).organization_id != org):
        raise HTTPException(422, "Parent Unit must belong to your organization")
    if payload.state_id and not db.get(State, payload.state_id):
        raise HTTPException(422, "State must be a configured master value")
    payload.other_industry_name = validate_industry_selection(db, payload.industry_type_id, payload.other_industry_name)
    contractor = Contractor(organization_id=org, **payload.model_dump())
    db.add(contractor); db.flush()
    generate_for_subject(db, org, "CONTRACTOR", contractor.id, date.today(), user.id)
    audit(db, organization_id=org, actor_id=user.id, action="CREATE_CONTRACTOR", module="ENTITIES", entity_type="Contractor", entity_id=contractor.id, new={"name": contractor.name})
    commit_or_conflict(db, "A contractor with this code already exists in the organization"); return entity_dict(db, contractor, "CONTRACTOR")


@router.patch("/contractors/{contractor_id}")
def update_contractor(contractor_id: str, payload: ContractorInput, user: User = Depends(current_user), db: Session = Depends(get_db)):
    contractor = db.get(Contractor, contractor_id)
    if not contractor or contractor.organization_id != org_for(user):
        raise HTTPException(404, "Contractor not found")
    require_org_admin_for_entity_management(db, user, contractor.organization_id)
    if db.scalar(select(Contractor).where(Contractor.organization_id == contractor.organization_id, Contractor.code == payload.code, Contractor.id != contractor.id)):
        raise HTTPException(409, "A contractor with this code already exists in the organization")
    if payload.unit_id and (not db.get(Unit, payload.unit_id) or db.get(Unit, payload.unit_id).organization_id != contractor.organization_id):
        raise HTTPException(422, "Parent Unit must belong to your organization")
    if payload.state_id and not db.get(State, payload.state_id):
        raise HTTPException(422, "State must be a configured master value")
    payload.other_industry_name = validate_industry_selection(db, payload.industry_type_id, payload.other_industry_name)
    for key, value in payload.model_dump().items(): setattr(contractor, key, value)
    generate_for_subject(db, contractor.organization_id, "CONTRACTOR", contractor.id, date.today(), user.id)
    audit(db, organization_id=contractor.organization_id, actor_id=user.id, action="UPDATE_CONTRACTOR", module="ENTITIES", entity_type="Contractor", entity_id=contractor.id, new={"name": contractor.name, "status": contractor.status})
    commit_or_conflict(db, "A contractor with this code already exists in the organization"); return entity_dict(db, contractor, "CONTRACTOR")


@router.delete("/contractors/{contractor_id}")
def delete_contractor(contractor_id: str, user: User = Depends(current_user), db: Session = Depends(get_db)):
    contractor = db.get(Contractor, contractor_id)
    if not contractor or contractor.organization_id != org_for(user):
        raise HTTPException(404, "Contractor not found")
    require_org_admin_for_entity_management(db, user, contractor.organization_id)
    require_deletable_entity(db, "CONTRACTOR", contractor.id)
    sites = list(db.scalars(select(ContractorSite).where(ContractorSite.contractor_id == contractor.id)))
    for site in sites:
        require_deletable_entity(db, "CONTRACTOR_SITE", site.id)
    for site in sites:
        db.delete(site)
    audit(db, organization_id=contractor.organization_id, actor_id=user.id, action="DELETE_CONTRACTOR", module="ENTITIES", entity_type="Contractor", entity_id=contractor.id, old={"name": contractor.name, "code": contractor.code})
    db.delete(contractor)
    db.commit()
    return {"ok": True}


@router.get("/contractor-sites")
def list_sites(paginated: bool = False, page: int = 1, page_size: int = 50, q: str | None = None, status: str | None = None, sort_by: str = "name", sort_dir: str = "asc", user: User = Depends(current_user), db: Session = Depends(get_db)):
    org = org_for(user)
    require_org_admin_for_entity_management(db, user, org)
    if paginated:
        return entity_page(db, ContractorSite, "CONTRACTOR_SITE", org, page, page_size, q, status, sort_by, sort_dir)
    return [entity_dict(db, row, "CONTRACTOR_SITE") for row in db.scalars(select(ContractorSite).where(ContractorSite.organization_id == org).order_by(ContractorSite.name))]


@router.get("/contractor-sites/{site_id}")
def get_site(site_id: str, user: User = Depends(current_user), db: Session = Depends(get_db)):
    site = db.get(ContractorSite, site_id)
    if not site or site.organization_id != org_for(user):
        raise HTTPException(404, "Contractor Site not found")
    require_org_admin_for_entity_management(db, user, site.organization_id)
    return entity_dict(db, site, "CONTRACTOR_SITE")


@router.post("/contractor-sites")
def create_site(payload: ContractorSiteInput, user: User = Depends(current_user), db: Session = Depends(get_db)):
    org = org_for(user); require_org_admin_for_entity_management(db, user, org)
    contractor = db.get(Contractor, payload.contractor_id)
    if not contractor or contractor.organization_id != org:
        raise HTTPException(422, "Contractor must belong to your organization")
    if db.scalar(select(ContractorSite).where(ContractorSite.contractor_id == contractor.id, ContractorSite.code == payload.code)):
        raise HTTPException(409, "A contractor site with this code already exists for the contractor")
    if payload.unit_id and (not db.get(Unit, payload.unit_id) or db.get(Unit, payload.unit_id).organization_id != org):
        raise HTTPException(422, "Linked Unit must belong to your organization")
    if not db.get(State, payload.state_id):
        raise HTTPException(422, "State must be a configured master value")
    payload.other_industry_name = validate_industry_selection(db, payload.industry_type_id, payload.other_industry_name)
    site = ContractorSite(organization_id=org, **payload.model_dump())
    db.add(site); db.flush(); generate_for_subject(db, org, "CONTRACTOR_SITE", site.id, date.today(), user.id)
    audit(db, organization_id=org, actor_id=user.id, action="CREATE_CONTRACTOR_SITE", module="ENTITIES", entity_type="ContractorSite", entity_id=site.id, new={"name": site.name})
    commit_or_conflict(db, "A contractor site with this code already exists for the contractor"); return entity_dict(db, site, "CONTRACTOR_SITE")


@router.patch("/contractor-sites/{site_id}")
def update_site(site_id: str, payload: ContractorSiteInput, user: User = Depends(current_user), db: Session = Depends(get_db)):
    site = db.get(ContractorSite, site_id)
    if not site or site.organization_id != org_for(user):
        raise HTTPException(404, "Contractor Site not found")
    require_org_admin_for_entity_management(db, user, site.organization_id)
    contractor = db.get(Contractor, payload.contractor_id)
    if not contractor or contractor.organization_id != site.organization_id:
        raise HTTPException(422, "Contractor must belong to your organization")
    if db.scalar(select(ContractorSite).where(ContractorSite.contractor_id == contractor.id, ContractorSite.code == payload.code, ContractorSite.id != site.id)):
        raise HTTPException(409, "A contractor site with this code already exists for the contractor")
    if not db.get(State, payload.state_id):
        raise HTTPException(422, "State must be a configured master value")
    payload.other_industry_name = validate_industry_selection(db, payload.industry_type_id, payload.other_industry_name)
    for key, value in payload.model_dump().items(): setattr(site, key, value)
    generate_for_subject(db, site.organization_id, "CONTRACTOR_SITE", site.id, date.today(), user.id)
    audit(db, organization_id=site.organization_id, actor_id=user.id, action="UPDATE_CONTRACTOR_SITE", module="ENTITIES", entity_type="ContractorSite", entity_id=site.id, new={"name": site.name, "status": site.status})
    commit_or_conflict(db, "A contractor site with this code already exists for the contractor"); return entity_dict(db, site, "CONTRACTOR_SITE")


@router.delete("/contractor-sites/{site_id}")
def delete_site(site_id: str, user: User = Depends(current_user), db: Session = Depends(get_db)):
    site = db.get(ContractorSite, site_id)
    if not site or site.organization_id != org_for(user):
        raise HTTPException(404, "Contractor Site not found")
    require_org_admin_for_entity_management(db, user, site.organization_id)
    require_deletable_entity(db, "CONTRACTOR_SITE", site.id)
    audit(db, organization_id=site.organization_id, actor_id=user.id, action="DELETE_CONTRACTOR_SITE", module="ENTITIES", entity_type="ContractorSite", entity_id=site.id, old={"name": site.name, "code": site.code})
    db.delete(site)
    db.commit()
    return {"ok": True}


@router.get("/users")
def list_users(paginated: bool = False, page: int = 1, page_size: int = 50, q: str | None = None, active: bool | None = None, sort_by: str = "name", sort_dir: str = "asc", user: User = Depends(current_user), db: Session = Depends(get_db)):
    org = org_for(user); require_org_admin(db, user, org)
    filters = [User.organization_id == org]
    if q:
        text = f"%{q.strip()}%"
        filters.append(or_(User.name.ilike(text), User.email.ilike(text)))
    if active is not None:
        filters.append(User.active.is_(active))
    allowed_sorts = {"name", "email", "active"}
    column = getattr(User, sort_by if sort_by in allowed_sorts else "name")
    ordering = column.desc() if sort_dir.lower() == "desc" else column.asc()
    if paginated:
        safe_page = max(page, 1)
        safe_size = min(max(page_size, 1), 100)
        total = db.scalar(select(func.count()).select_from(User).where(*filters)) or 0
        rows = db.scalars(select(User).where(*filters).order_by(ordering, User.id).offset((safe_page - 1) * safe_size).limit(safe_size)).all()
        return {
            "items": [user_summary(db, row) | {"active": row.active, "mobile": row.mobile} for row in rows],
            "total": total,
            "page": safe_page,
            "page_size": safe_size,
        }
    return [user_summary(db, row) | {"active": row.active} for row in db.scalars(select(User).where(*filters).order_by(ordering, User.id))]


@router.post("/users")
def create_user(payload: UserInput, user: User = Depends(current_user), db: Session = Depends(get_db)):
    org = org_for(user); require_org_admin(db, user, org)
    if db.scalar(select(User).where(User.email == payload.email.lower())):
        raise HTTPException(409, "A user with this email already exists")
    new_user = User(organization_id=org, name=payload.name, email=payload.email.lower(), mobile=payload.mobile, password_hash=hash_password(payload.password), active=payload.active)
    db.add(new_user); db.flush()
    audit(db, organization_id=org, actor_id=user.id, action="CREATE_USER", module="USERS", entity_type="User", entity_id=new_user.id, new={"email": new_user.email})
    db.commit(); return user_summary(db, new_user)


@router.get("/users/{user_id}")
def get_user(user_id: str, user: User = Depends(current_user), db: Session = Depends(get_db)):
    org = org_for(user); require_org_admin(db, user, org)
    target = db.get(User, user_id)
    if not target or target.organization_id != org: raise HTTPException(404, "User not found")
    return user_summary(db, target) | {"active": target.active, "mobile": target.mobile}


@router.patch("/users/{user_id}")
def update_user(user_id: str, payload: UserUpdateInput, user: User = Depends(current_user), db: Session = Depends(get_db)):
    org = org_for(user); require_org_admin(db, user, org)
    target = db.get(User, user_id)
    if not target or target.organization_id != org: raise HTTPException(404, "User not found")
    old = {"name": target.name, "active": target.active, "mobile": target.mobile}
    for key, value in payload.model_dump(exclude_unset=True).items(): setattr(target, key, value)
    audit(db, organization_id=org, actor_id=user.id, action="UPDATE_USER", module="USERS", entity_type="User", entity_id=target.id, old=old, new={"name": target.name, "active": target.active, "mobile": target.mobile})
    db.commit(); return user_summary(db, target) | {"active": target.active, "mobile": target.mobile}


@router.put("/users/{user_id}/role-scopes")
def replace_role_scopes(user_id: str, payload: list[RoleScopeInput], user: User = Depends(current_user), db: Session = Depends(get_db)):
    org = org_for(user); require_org_admin(db, user, org)
    target = db.get(User, user_id)
    if not target or target.organization_id != org: raise HTTPException(404, "User not found")
    valid_roles = {"ORGANIZATION_ADMIN", "UNIT_ADMIN", "UNIT_MAKER", "UNIT_CHECKER", "CONTRACTOR_ADMIN", "CONTRACTOR_MAKER", "CONTRACTOR_CHECKER", "VIEWER", "AUDITOR"}
    if any(item.role not in valid_roles for item in payload): raise HTTPException(422, "Unsupported MVP role")
    seen_scopes: set[tuple[str, str, str]] = set()
    for item in payload:
        key = (item.role, item.scope_type, item.scope_id)
        if key in seen_scopes:
            raise HTTPException(422, "Duplicate role scope is not allowed")
        seen_scopes.add(key)
        require_scope_belongs_to_organization(db, org, item.scope_type, item.scope_id)
    db.query(UserRoleScope).filter_by(user_id=user_id).delete()
    db.add_all([UserRoleScope(user_id=user_id, **item.model_dump()) for item in payload])
    audit(db, organization_id=org, actor_id=user.id, action="UPDATE_ROLE_SCOPES", module="USERS", entity_type="User", entity_id=user_id, new=[item.model_dump() for item in payload])
    commit_or_conflict(db, "Role scope could not be saved because it conflicts with an existing scope"); return user_summary(db, target)


@router.get("/compliance-master/template")
def compliance_template(user: User = Depends(current_user), db: Session = Depends(get_db)):
    require_org_admin(db, user, org_for(user))
    headers = ["Compliance ID", "Compliance Name", "Version", "Entity Type", "State", "Industry Type", "Frequency", "Due Date Rule", "Due Date Offset", "Due Date Anchor", "Required Document", "Risk Level", "Effective From", "Effective To", "Active"]
    workbook = Workbook(); worksheet = workbook.active; worksheet.title = "Compliance Master"; worksheet.append(headers)
    worksheet.append(["DEMO-MONTHLY-001", "Monthly demonstration filing", 1, "UNIT", "Tamil Nadu", "Factory", "MONTHLY", "FIXED_DAY_OF_MONTH", 15, "", "Monthly Return", "MEDIUM", "2026-01-01", "", "TRUE"])
    stream = io.BytesIO(); workbook.save(stream); stream.seek(0)
    return StreamingResponse(stream, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", headers={"Content-Disposition": "attachment; filename=compliance-master-template.xlsx"})


def rule_version_dict(version: ComplianceRuleVersion, rule: ComplianceRule, criteria: RuleApplicabilityCriteria) -> dict:
    return {
        "id": version.id, "rule_id": rule.id, "compliance_id": rule.compliance_id,
        "name": version.name, "description": version.description, "act": version.act,
        "rule_reference": version.rule_reference, "section": version.section,
        "compliance_type": version.compliance_type, "document_type": version.document_type,
        "form_number": version.form_number, "legal_description": version.legal_description,
        "consequence_or_penalty": version.consequence_or_penalty, "version": version.version,
        "frequency": version.frequency, "entity_type": criteria.entity_type,
        "state_id": criteria.state_id, "industry_type_id": criteria.industry_type_id,
        "due_date_rule": version.due_date_rule, "due_date_offset": version.due_date_offset,
        "due_date_anchor": version.due_date_anchor, "grace_days": version.grace_days,
        "required_document": version.required_document, "risk_level": version.risk_level,
        "effective_from": version.effective_from, "effective_to": version.effective_to,
        "active": version.active, "created_at": version.created_at, "updated_at": version.updated_at,
    }


@router.get("/compliance-rules")
def list_rules(paginated: bool = False, page: int = 1, page_size: int = 50, q: str | None = None, frequency: str | None = None, active: bool | None = None, sort_by: str = "compliance_id", sort_dir: str = "asc", user: User = Depends(current_user), db: Session = Depends(get_db)):
    org = org_for(user)
    require_org_admin(db, user, org)
    query = db.query(ComplianceRuleVersion, ComplianceRule, RuleApplicabilityCriteria).join(ComplianceRule, ComplianceRule.id == ComplianceRuleVersion.rule_id).join(RuleApplicabilityCriteria, RuleApplicabilityCriteria.rule_version_id == ComplianceRuleVersion.id).filter(ComplianceRule.organization_id == org)
    if q:
        text = f"%{q.strip()}%"
        query = query.filter(or_(ComplianceRule.compliance_id.ilike(text), ComplianceRuleVersion.name.ilike(text), ComplianceRuleVersion.act.ilike(text)))
    if frequency:
        query = query.filter(ComplianceRuleVersion.frequency == frequency.upper())
    if active is not None:
        query = query.filter(ComplianceRuleVersion.active.is_(active))
    ordering = {
        "compliance_id": ComplianceRule.compliance_id,
        "name": ComplianceRuleVersion.name,
        "frequency": ComplianceRuleVersion.frequency,
        "risk_level": ComplianceRuleVersion.risk_level,
        "effective_from": ComplianceRuleVersion.effective_from,
        "active": ComplianceRuleVersion.active,
    }.get(sort_by, ComplianceRule.compliance_id)
    order = ordering.desc() if sort_dir.lower() == "desc" else ordering.asc()
    query = query.order_by(order, ComplianceRuleVersion.version.desc())
    if not paginated:
        return [rule_version_dict(version, rule, criteria) for version, rule, criteria in query.all()]
    safe_page = max(page, 1)
    safe_size = min(max(page_size, 1), 100)
    total = query.count()
    rows = query.offset((safe_page - 1) * safe_size).limit(safe_size).all()
    return {"items": [rule_version_dict(version, rule, criteria) for version, rule, criteria in rows], "total": total, "page": safe_page, "page_size": safe_size}


@router.get("/compliance-rules/{version_id}")
def get_rule_version(version_id: str, user: User = Depends(current_user), db: Session = Depends(get_db)):
    org = org_for(user)
    require_org_admin(db, user, org)
    row = db.query(ComplianceRuleVersion, ComplianceRule, RuleApplicabilityCriteria).join(ComplianceRule, ComplianceRule.id == ComplianceRuleVersion.rule_id).join(RuleApplicabilityCriteria, RuleApplicabilityCriteria.rule_version_id == ComplianceRuleVersion.id).filter(ComplianceRuleVersion.id == version_id, ComplianceRule.organization_id == org).first()
    if not row:
        raise HTTPException(404, "Compliance rule version not found")
    return rule_version_dict(*row)


def create_rule_version(db: Session, org: str, actor: User, payload: RuleInput):
    if payload.entity_type not in SUBJECTS: raise HTTPException(422, "Entity Type must be UNIT, CONTRACTOR, or CONTRACTOR_SITE")
    rule = db.scalar(select(ComplianceRule).where(ComplianceRule.organization_id == org, ComplianceRule.compliance_id == payload.compliance_id))
    if not rule:
        if payload.version != 1: raise HTTPException(422, "New rules must begin with version 1")
        rule = ComplianceRule(organization_id=org, compliance_id=payload.compliance_id, created_by_id=actor.id); db.add(rule); db.flush()
    elif db.scalar(select(ComplianceRuleVersion).where(ComplianceRuleVersion.rule_id == rule.id, ComplianceRuleVersion.version == payload.version)):
        raise HTTPException(409, "This rule version already exists")
    elif payload.version != (db.scalar(select(func.max(ComplianceRuleVersion.version)).where(ComplianceRuleVersion.rule_id == rule.id)) or 0) + 1:
        raise HTTPException(422, "Rule version must be the next sequential version")
    version = ComplianceRuleVersion(rule_id=rule.id, version=payload.version, name=payload.name, description=payload.description, act=payload.act, rule_reference=payload.rule_reference, section=payload.section, compliance_type=payload.compliance_type, document_type=payload.document_type, form_number=payload.form_number, legal_description=payload.legal_description, consequence_or_penalty=payload.consequence_or_penalty, frequency=payload.frequency.upper(), due_date_rule=payload.due_date_rule.upper(), due_date_offset=payload.due_date_offset, due_date_anchor=payload.due_date_anchor, grace_days=payload.grace_days, required_document=payload.required_document, risk_level=payload.risk_level.upper(), effective_from=payload.effective_from, effective_to=payload.effective_to, active=payload.active)
    db.add(version); db.flush(); db.add(RuleApplicabilityCriteria(rule_version_id=version.id, entity_type=payload.entity_type, state_id=payload.state_id, industry_type_id=payload.industry_type_id))
    audit(db, organization_id=org, actor_id=actor.id, action="CREATE_RULE_VERSION", module="COMPLIANCE_MASTER", entity_type="ComplianceRuleVersion", entity_id=version.id, new={"compliance_id": payload.compliance_id, "version": payload.version})
    return version


@router.post("/compliance-rules")
def create_rule(payload: RuleInput, user: User = Depends(current_user), db: Session = Depends(get_db)):
    org = org_for(user); require_org_admin(db, user, org)
    version = create_rule_version(db, org, user, payload); db.commit()
    return {"id": version.id, "message": "Rule version created"}


@router.patch("/compliance-rules/{version_id}/status")
def set_rule_version_status(version_id: str, payload: RuleStatusInput, user: User = Depends(current_user), db: Session = Depends(get_db)):
    org = org_for(user)
    require_org_admin(db, user, org)
    row = db.query(ComplianceRuleVersion, ComplianceRule, RuleApplicabilityCriteria).join(ComplianceRule, ComplianceRule.id == ComplianceRuleVersion.rule_id).join(RuleApplicabilityCriteria, RuleApplicabilityCriteria.rule_version_id == ComplianceRuleVersion.id).filter(ComplianceRuleVersion.id == version_id, ComplianceRule.organization_id == org).first()
    if not row:
        raise HTTPException(404, "Compliance rule version not found")
    version, rule, criteria = row
    old = {"active": version.active}
    version.active = payload.active
    audit(db, organization_id=org, actor_id=user.id, action="ACTIVATE_RULE_VERSION" if payload.active else "DEACTIVATE_RULE_VERSION", module="COMPLIANCE_MASTER", entity_type="ComplianceRuleVersion", entity_id=version.id, old=old, new={"active": version.active, "compliance_id": rule.compliance_id, "version": version.version})
    db.commit()
    return rule_version_dict(version, rule, criteria)


@router.post("/compliance-imports/validate")
async def validate_import(file: UploadFile = File(...), user: User = Depends(current_user), db: Session = Depends(get_db)):
    org = org_for(user); require_org_admin(db, user, org)
    if not (file.filename or "").lower().endswith(".xlsx"): raise HTTPException(422, "Compliance Master must be an .xlsx file")
    content = await file.read()
    try: workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    except Exception as exc: raise HTTPException(422, "The uploaded file is not a valid Excel workbook") from exc
    if "COW" in workbook.sheetnames:
        return stage_construction_checklist(workbook, org, user, db)
    if "Compliance Master" not in workbook.sheetnames: raise HTTPException(422, "Worksheet 'Compliance Master' is required")
    sheet = workbook["Compliance Master"]
    headers = [cell.value for cell in next(sheet.iter_rows(min_row=1, max_row=1))]
    required = {"Compliance ID", "Compliance Name", "Version", "Entity Type", "Frequency", "Due Date Rule", "Effective From"}
    if not required.issubset(set(headers)): raise HTTPException(422, f"Missing required headers: {', '.join(sorted(required - set(headers)))}")
    import_record = ComplianceImport(organization_id=org, uploaded_by_id=user.id, original_filename=file.filename or "compliance-master.xlsx")
    db.add(import_record); db.flush(); seen = set(); valid = errors = total = 0
    states_by_name = {state.name.lower(): state.id for state in db.scalars(select(State))}; industries_by_name = {industry.name.lower(): industry.id for industry in db.scalars(select(IndustryType))}
    for number, values in enumerate(sheet.iter_rows(min_row=2, values_only=True), start=2):
        if not any(value is not None and str(value).strip() for value in values): continue
        total += 1; row = {headers[index]: values[index] for index in range(len(headers))}; messages = []
        for field in required:
            if row.get(field) in (None, ""): messages.append((field, "Value is required"))
        identity = (str(row.get("Compliance ID", "")).strip(), str(row.get("Version", "")).strip())
        if identity in seen: messages.append(("Compliance ID", "Duplicate Compliance ID and Version in uploaded file"))
        seen.add(identity)
        if str(row.get("Entity Type", "")).upper() not in SUBJECTS: messages.append(("Entity Type", "Must be UNIT, CONTRACTOR, or CONTRACTOR_SITE"))
        if str(row.get("Frequency", "")).upper() not in {"ONE_TIME", "MONTHLY", "QUARTERLY", "HALF_YEARLY", "ANNUAL", "EVENT_BASED"}: messages.append(("Frequency", "Unsupported MVP frequency"))
        if row.get("State") and str(row["State"]).strip().lower() not in states_by_name: messages.append(("State", "State is not configured"))
        if row.get("Industry Type") and str(row["Industry Type"]).strip().lower() not in industries_by_name: messages.append(("Industry Type", "Industry Type is not configured"))
        try: int(row.get("Version")); date.fromisoformat(str(row.get("Effective From")))
        except (TypeError, ValueError): messages.append(("Version/Effective From", "Version must be an integer and Effective From must be ISO date"))
        import_row = ComplianceImportRow(import_id=import_record.id, row_number=number, payload_json=json.dumps(row, default=str), valid=not messages); db.add(import_row); db.flush()
        for field, message in messages: db.add(ComplianceImportError(import_row_id=import_row.id, field=field, message=message))
        if messages: errors += 1
        else: valid += 1
    import_record.total_rows = total; import_record.valid_rows = valid; import_record.error_rows = errors
    audit(db, organization_id=org, actor_id=user.id, action="VALIDATE_IMPORT", module="COMPLIANCE_MASTER", entity_type="ComplianceImport", entity_id=import_record.id, new={"total": total, "valid": valid, "errors": errors})
    db.commit(); return {"id": import_record.id, "total_rows": total, "valid_rows": valid, "error_rows": errors}


@router.get("/compliance-imports/{import_id}/preview")
def import_preview(import_id: str, user: User = Depends(current_user), db: Session = Depends(get_db)):
    record = db.get(ComplianceImport, import_id); org = org_for(user)
    if not record or record.organization_id != org: raise HTTPException(404, "Import not found")
    require_org_admin(db, user, org)
    rows = []
    for item in db.scalars(select(ComplianceImportRow).where(ComplianceImportRow.import_id == import_id)):
        errors = db.query(ComplianceImportError).filter_by(import_row_id=item.id).all()
        rows.append({"row_number": item.row_number, "valid": item.valid, "data": json.loads(item.payload_json), "errors": [{"field": error.field, "message": error.message} for error in errors]})
    return {"import": model_dict(record, [field.name for field in record.__table__.columns]), "rows": rows}


@router.post("/compliance-imports/{import_id}/confirm")
def confirm_import(import_id: str, user: User = Depends(current_user), db: Session = Depends(get_db)):
    org = org_for(user); require_org_admin(db, user, org); record = db.get(ComplianceImport, import_id)
    if not record or record.organization_id != org: raise HTTPException(404, "Import not found")
    if record.status != "VALIDATED": raise HTTPException(409, "Import has already been confirmed or cancelled")
    state_names = {state.name.lower(): state.id for state in db.scalars(select(State))}; industry_names = {industry.name.lower(): industry.id for industry in db.scalars(select(IndustryType))}
    imported = skipped = 0
    for item in db.scalars(select(ComplianceImportRow).where(ComplianceImportRow.import_id == import_id, ComplianceImportRow.valid.is_(True))):
        row = json.loads(item.payload_json)
        payload = RuleInput(compliance_id=str(row["Compliance ID"]).strip(), name=str(row["Compliance Name"]).strip(), description=str(row["Description"]) if row.get("Description") else None, act=str(row["Act"]) if row.get("Act") else None, rule_reference=str(row["Rule"]) if row.get("Rule") else None, section=str(row["Section"]) if row.get("Section") else None, compliance_type=str(row["Compliance Type"]) if row.get("Compliance Type") else None, document_type=str(row["Document Type"]) if row.get("Document Type") else None, form_number=str(row["Form Number"]) if row.get("Form Number") else None, legal_description=str(row["Legal Description"]) if row.get("Legal Description") else None, consequence_or_penalty=str(row["Consequence / Penalty"]) if row.get("Consequence / Penalty") else None, version=int(row["Version"]), entity_type=str(row["Entity Type"]).upper(), state_id=state_names.get(str(row.get("State") or "").lower()), industry_type_id=industry_names.get(str(row.get("Industry Type") or "").lower()), frequency=str(row["Frequency"]).upper(), due_date_rule=str(row["Due Date Rule"]).upper(), due_date_offset=int(row["Due Date Offset"]) if row.get("Due Date Offset") not in (None, "") else None, due_date_anchor=str(row["Due Date Anchor"]) if row.get("Due Date Anchor") else None, required_document=str(row["Required Document"]) if row.get("Required Document") else None, risk_level=str(row.get("Risk Level") or "MEDIUM").upper(), effective_from=date.fromisoformat(str(row["Effective From"])), effective_to=date.fromisoformat(str(row["Effective To"])) if row.get("Effective To") else None, active=parse_bool(row.get("Active")))
        existing_rule = db.scalar(select(ComplianceRule).where(ComplianceRule.organization_id == org, ComplianceRule.compliance_id == payload.compliance_id))
        if existing_rule and db.scalar(select(ComplianceRuleVersion).where(ComplianceRuleVersion.rule_id == existing_rule.id, ComplianceRuleVersion.version == payload.version)):
            skipped += 1
            continue
        create_rule_version(db, org, user, payload); imported += 1
    record.status = "CONFIRMED"; record.confirmed_at = datetime.now(UTC)
    audit(db, organization_id=org, actor_id=user.id, action="CONFIRM_IMPORT", module="COMPLIANCE_MASTER", entity_type="ComplianceImport", entity_id=record.id, new={"imported": imported, "skipped": skipped})
    db.commit(); return {"id": record.id, "imported": imported, "skipped": skipped, "error_rows": record.error_rows}


@router.get("/compliance-imports")
def list_imports(user: User = Depends(current_user), db: Session = Depends(get_db)):
    org = org_for(user); require_org_admin(db, user, org)
    return [model_dict(item, [field.name for field in item.__table__.columns]) for item in db.scalars(select(ComplianceImport).where(ComplianceImport.organization_id == org).order_by(ComplianceImport.created_at.desc()))]


@router.post("/compliance-generation/run")
def run_generation(payload: GenerationRequest, user: User = Depends(current_user), db: Session = Depends(get_db)):
    org = org_for(user); require_manager(db, user, org); as_of = payload.as_of_date or date.today(); created = []
    subject_types = [payload.subject_type] if payload.subject_type else list(SUBJECTS)
    for subject_type in subject_types:
        model = SUBJECTS.get(subject_type)
        if not model: raise HTTPException(422, "Unsupported subject type")
        rows = [db.get(model, payload.subject_id)] if payload.subject_id else list(db.scalars(select(model).where(model.organization_id == org)))
        for subject in rows:
            if subject and subject.organization_id == org and can_access_subject(db, user, org, subject_type, subject.id): created.extend(generate_for_subject(db, org, subject_type, subject.id, as_of, user.id))
    db.commit(); return {"created": len(created), "instance_ids": [item.id for item in created]}


@router.get("/compliance-instances")
def list_instances(frequency: str | None = None, status: str | None = None, risk_level: str | None = None, entity_type: str | None = None, date_from: date | None = None, date_to: date | None = None, q: str | None = None, page: int = 1, page_size: int = 50, user: User = Depends(current_user), db: Session = Depends(get_db)):
    org = org_for(user); instances = list(db.scalars(select(ComplianceInstance).where(ComplianceInstance.organization_id == org).order_by(ComplianceInstance.due_date)))
    results = []
    for instance in instances:
        if not can_access_subject(db, user, org, instance.subject_type, instance.subject_id): continue
        row = instance_dict(db, instance)
        if frequency and row["frequency"] != frequency: continue
        if status and status not in {instance.status, row["display_status"]}: continue
        if risk_level and row["risk_level"] != risk_level: continue
        if entity_type and instance.subject_type != entity_type: continue
        if date_from and instance.due_date < date_from: continue
        if date_to and instance.due_date > date_to: continue
        if q and q.lower() not in " ".join(str(row.get(key) or "") for key in ("compliance_id", "compliance_name", "subject_name")).lower(): continue
        results.append(row)
    return list_page(results, page, min(page_size, 100))


def get_instance_or_404(db: Session, user: User, instance_id: str) -> ComplianceInstance:
    instance = db.get(ComplianceInstance, instance_id)
    if not instance or instance.organization_id != org_for(user) or not can_access_subject(db, user, instance.organization_id, instance.subject_type, instance.subject_id): raise HTTPException(404, "Compliance instance not found")
    return instance


@router.get("/compliance-instances/{instance_id}")
def get_instance(instance_id: str, user: User = Depends(current_user), db: Session = Depends(get_db)):
    instance = get_instance_or_404(db, user, instance_id); result = instance_dict(db, instance, detail=True)
    result["history"] = [model_dict(item, [field.name for field in item.__table__.columns]) for item in db.scalars(select(ComplianceWorkflowHistory).where(ComplianceWorkflowHistory.instance_id == instance.id).order_by(ComplianceWorkflowHistory.created_at))]
    result["comments"] = [model_dict(item, [field.name for field in item.__table__.columns]) for item in db.scalars(select(ComplianceComment).where(ComplianceComment.instance_id == instance.id).order_by(ComplianceComment.created_at))]
    return result


@router.patch("/compliance-instances/{instance_id}/activity")
def save_activity(instance_id: str, payload: ActivityInput, user: User = Depends(current_user), db: Session = Depends(get_db)):
    instance = get_instance_or_404(db, user, instance_id); roles = roles_for(db, user.id)
    if not roles & MAKER_ROLES: raise HTTPException(403, "Maker permission is required")
    assert_assignment_or_admin(db, instance, user.id, roles, "MAKER")
    if instance.row_version != payload.row_version: raise HTTPException(409, "This compliance was changed by another user; refresh and try again")
    if instance.status in {"APPROVED", "SUBMITTED", "UNDER_REVIEW"}: raise HTTPException(409, "This compliance is not editable in its current state")
    old = {"status": instance.status}; instance.activity_reference = payload.filing_reference; instance.completed_on = payload.completed_on; instance.amount = payload.amount; instance.remarks = payload.remarks
    if instance.status == "PENDING": transition(db, instance, user.id, "START_WORK", "IN_PROGRESS")
    else: instance.row_version += 1
    audit(db, organization_id=instance.organization_id, actor_id=user.id, action="SAVE_COMPLIANCE_ACTIVITY", module="COMPLIANCE", entity_type="ComplianceInstance", entity_id=instance.id, old=old, new={"status": instance.status})
    db.commit(); return instance_dict(db, instance, detail=True)


@router.post("/compliance-instances/{instance_id}/assignments")
def assign(instance_id: str, payload: AssignmentInput, user: User = Depends(current_user), db: Session = Depends(get_db)):
    instance = get_instance_or_404(db, user, instance_id); require_manager(db, user, instance.organization_id)
    if payload.assignment_type not in {"MAKER", "CHECKER"}: raise HTTPException(422, "Assignment type must be MAKER or CHECKER")
    assignee = db.get(User, payload.user_id)
    if not assignee or assignee.organization_id != instance.organization_id or not assignee.active: raise HTTPException(422, "Assignee must be an active user in the organization")
    for prior in db.query(ComplianceAssignment).filter_by(instance_id=instance.id, assignment_type=payload.assignment_type, active=True): prior.active = False; prior.ended_at = datetime.now(UTC)
    assignment = ComplianceAssignment(instance_id=instance.id, user_id=assignee.id, assignment_type=payload.assignment_type, assigned_by_id=user.id); db.add(assignment)
    db.add(ComplianceWorkflowHistory(instance_id=instance.id, actor_id=user.id, action="ASSIGN", from_status=instance.status, to_status=instance.status, comment=f"{payload.assignment_type}: {assignee.name}"))
    audit(db, organization_id=instance.organization_id, actor_id=user.id, action="ASSIGN_COMPLIANCE", module="COMPLIANCE", entity_type="ComplianceInstance", entity_id=instance.id, new={"assignment_type": payload.assignment_type, "user_id": assignee.id})
    db.commit(); return model_dict(assignment, [field.name for field in assignment.__table__.columns])


@router.post("/compliance-instances/{instance_id}/evidence")
async def upload_evidence(instance_id: str, category: str = Query("SUPPORTING_DOCUMENT"), file: UploadFile = File(...), user: User = Depends(current_user), db: Session = Depends(get_db)):
    instance = get_instance_or_404(db, user, instance_id); roles = roles_for(db, user.id)
    if not roles & MAKER_ROLES: raise HTTPException(403, "Maker permission is required")
    assert_assignment_or_admin(db, instance, user.id, roles, "MAKER")
    if instance.status in {"APPROVED", "SUBMITTED", "UNDER_REVIEW"}: raise HTTPException(409, "Evidence cannot be changed in this workflow state")
    stored_name, size, checksum = await save_upload(file); version = (db.scalar(select(func.max(ComplianceEvidence.version)).where(ComplianceEvidence.instance_id == instance.id, ComplianceEvidence.category == category)) or 0) + 1
    evidence = ComplianceEvidence(instance_id=instance.id, category=category, version=version, original_filename=file.filename or "evidence", stored_filename=stored_name, content_type=file.content_type, size_bytes=size, checksum_sha256=checksum, uploaded_by_id=user.id)
    db.add(evidence); db.flush()
    audit(db, organization_id=instance.organization_id, actor_id=user.id, action="UPLOAD_EVIDENCE", module="EVIDENCE", entity_type="ComplianceEvidence", entity_id=evidence.id, new={"filename": evidence.original_filename, "version": version})
    try:
        db.commit()
    except Exception:
        db.rollback()
        delete_stored_file(stored_name)
        raise
    return model_dict(evidence, [field.name for field in evidence.__table__.columns])


@router.get("/evidence/{evidence_id}/download")
def download_evidence(evidence_id: str, user: User = Depends(current_user), db: Session = Depends(get_db)):
    evidence = db.get(ComplianceEvidence, evidence_id)
    if not evidence: raise HTTPException(404, "Evidence not found")
    instance = get_instance_or_404(db, user, evidence.instance_id)
    return FileResponse(file_path(evidence.stored_filename), filename=evidence.original_filename, media_type=evidence.content_type or "application/octet-stream")


@router.post("/evidence/{evidence_id}/verify")
def verify_evidence(evidence_id: str, payload: EvidenceVerificationInput, user: User = Depends(current_user), db: Session = Depends(get_db)):
    evidence = db.get(ComplianceEvidence, evidence_id)
    if not evidence: raise HTTPException(404, "Evidence not found")
    instance = get_instance_or_404(db, user, evidence.instance_id)
    roles = roles_for(db, user.id)
    if not roles & CHECKER_ROLES: raise HTTPException(403, "Checker permission is required")
    assert_assignment_or_admin(db, instance, user.id, roles, "CHECKER")
    state = payload.verification_state.upper()
    if state not in {"VERIFIED", "REJECTED"}: raise HTTPException(422, "Verification state must be VERIFIED or REJECTED")
    if state == "REJECTED" and not payload.reason: raise HTTPException(422, "A rejection reason is required")
    evidence.verification_state = state; evidence.verification_reason = payload.reason; evidence.verified_by_id = user.id; evidence.verified_at = datetime.now(UTC)
    audit(db, organization_id=instance.organization_id, actor_id=user.id, action=f"{state}_EVIDENCE", module="EVIDENCE", entity_type="ComplianceEvidence", entity_id=evidence.id, new={"state": state}, reason=payload.reason)
    db.commit(); return model_dict(evidence, [field.name for field in evidence.__table__.columns])


@router.post("/compliance-instances/{instance_id}/not-applicable")
def mark_not_applicable(instance_id: str, payload: NotApplicableInput, user: User = Depends(current_user), db: Session = Depends(get_db)):
    instance = get_instance_or_404(db, user, instance_id)
    require_manager(db, user, instance.organization_id)
    if instance.status in {"APPROVED", "NOT_APPLICABLE"}: raise HTTPException(409, "Completed items cannot be marked not applicable")
    instance.not_applicable_reason = payload.reason
    transition(db, instance, user.id, "MARK_NOT_APPLICABLE", "NOT_APPLICABLE", payload.reason)
    db.commit(); return instance_dict(db, instance, detail=True)


@router.post("/compliance-instances/{instance_id}/submit")
def submit_instance(instance_id: str, user: User = Depends(current_user), db: Session = Depends(get_db)):
    instance = get_instance_or_404(db, user, instance_id); roles = roles_for(db, user.id)
    if not roles & MAKER_ROLES: raise HTTPException(403, "Maker permission is required")
    assert_assignment_or_admin(db, instance, user.id, roles, "MAKER"); submit(db, instance, user.id)
    for assignment in db.query(ComplianceAssignment).filter_by(instance_id=instance.id, assignment_type="CHECKER", active=True):
        db.add(InAppNotification(organization_id=instance.organization_id, recipient_id=assignment.user_id, notification_type="PENDING_APPROVAL", title="Compliance awaiting your review", body=instance.activity_reference, reference_type="ComplianceInstance", reference_id=instance.id))
    db.commit(); return instance_dict(db, instance, detail=True)


@router.post("/compliance-instances/{instance_id}/begin-review")
def begin_review(instance_id: str, payload: WorkflowDecision, user: User = Depends(current_user), db: Session = Depends(get_db)):
    instance = get_instance_or_404(db, user, instance_id); roles = roles_for(db, user.id)
    if not roles & CHECKER_ROLES: raise HTTPException(403, "Checker permission is required")
    assert_assignment_or_admin(db, instance, user.id, roles, "CHECKER")
    if instance.status != "SUBMITTED": raise HTTPException(409, "Only submitted compliances can enter review")
    if payload.row_version != instance.row_version: raise HTTPException(409, "This compliance was changed by another user")
    transition(db, instance, user.id, "BEGIN_REVIEW", "UNDER_REVIEW", payload.comment); db.commit(); return instance_dict(db, instance, detail=True)


def reviewer_action(instance_id: str, payload: WorkflowDecision, target: str, action: str, user: User, db: Session):
    instance = get_instance_or_404(db, user, instance_id); roles = roles_for(db, user.id)
    if not roles & CHECKER_ROLES: raise HTTPException(403, "Checker permission is required")
    assert_assignment_or_admin(db, instance, user.id, roles, "CHECKER")
    if instance.status != "UNDER_REVIEW": raise HTTPException(409, "Compliance must be under review")
    if payload.row_version != instance.row_version: raise HTTPException(409, "This compliance was changed by another user")
    latest_submission = db.scalar(select(ComplianceSubmission).where(ComplianceSubmission.instance_id == instance.id).order_by(ComplianceSubmission.revision.desc()))
    if latest_submission and latest_submission.submitted_by_id == user.id:
        raise HTTPException(403, "A maker cannot approve or reject their own submission")
    if target == "APPROVED":
        version = db.get(ComplianceRuleVersion, instance.rule_version_id)
        if version.required_document and not db.query(ComplianceEvidence).filter_by(instance_id=instance.id, verification_state="VERIFIED").first():
            raise HTTPException(422, "Required evidence must be verified before approval")
    transition(db, instance, user.id, action, target, payload.comment)
    if target in {"CORRECTION_REQUIRED", "REJECTED"}:
        for assignment in db.query(ComplianceAssignment).filter_by(instance_id=instance.id, assignment_type="MAKER", active=True):
            db.add(InAppNotification(organization_id=instance.organization_id, recipient_id=assignment.user_id, notification_type=target, title="Compliance requires your attention", body=payload.comment, reference_type="ComplianceInstance", reference_id=instance.id))
    if target == "APPROVED":
        instance.approved_at = datetime.now(UTC); instance.approved_by_id = user.id; instance.completion_late = instance.approved_at.date() > instance.due_date
    db.commit(); return instance_dict(db, instance, detail=True)


@router.post("/compliance-instances/{instance_id}/approve")
def approve(instance_id: str, payload: WorkflowDecision, user: User = Depends(current_user), db: Session = Depends(get_db)):
    return reviewer_action(instance_id, payload, "APPROVED", "APPROVE", user, db)


@router.post("/compliance-instances/{instance_id}/reject")
def reject(instance_id: str, payload: WorkflowDecision, user: User = Depends(current_user), db: Session = Depends(get_db)):
    return reviewer_action(instance_id, payload, "REJECTED", "REJECT", user, db)


@router.post("/compliance-instances/{instance_id}/request-correction")
def correction(instance_id: str, payload: WorkflowDecision, user: User = Depends(current_user), db: Session = Depends(get_db)):
    return reviewer_action(instance_id, payload, "CORRECTION_REQUIRED", "REQUEST_CORRECTION", user, db)


@router.post("/compliance-instances/{instance_id}/comments")
def add_comment(instance_id: str, payload: CommentInput, user: User = Depends(current_user), db: Session = Depends(get_db)):
    instance = get_instance_or_404(db, user, instance_id); comment = ComplianceComment(instance_id=instance.id, author_id=user.id, body=payload.body); db.add(comment); db.flush()
    audit(db, organization_id=instance.organization_id, actor_id=user.id, action="ADD_COMMENT", module="COMPLIANCE", entity_type="ComplianceComment", entity_id=comment.id); db.commit(); return model_dict(comment, [field.name for field in comment.__table__.columns])


@router.get("/documents")
def list_documents(q: str | None = None, entity_type: str | None = None, verification_state: str | None = None, paginated: bool = False, page: int = 1, page_size: int = 50, sort_by: str = "created_at", sort_dir: str = "desc", user: User = Depends(current_user), db: Session = Depends(get_db)):
    org = org_for(user)
    documents = []
    for evidence, instance in db.query(ComplianceEvidence, ComplianceInstance).join(ComplianceInstance, ComplianceInstance.id == ComplianceEvidence.instance_id).filter(ComplianceInstance.organization_id == org).order_by(ComplianceEvidence.created_at.desc()):
        if entity_type and instance.subject_type != entity_type:
            continue
        if verification_state and evidence.verification_state != verification_state.upper():
            continue
        if not can_access_subject(db, user, org, instance.subject_type, instance.subject_id):
            continue
        if q and q.lower() not in evidence.original_filename.lower():
            continue
        row = model_dict(evidence, [field.name for field in evidence.__table__.columns])
        detail = instance_dict(db, instance)
        row.update({"compliance_instance_id": instance.id, "entity_type": instance.subject_type, "entity_id": instance.subject_id, "period": instance.period_key, "subject_name": detail["subject_name"], "compliance_name": detail["compliance_name"]})
        documents.append(row)
    sort_key = {
        "created_at": lambda item: str(item.get("created_at") or ""),
        "filename": lambda item: str(item.get("original_filename") or "").lower(),
        "verification_state": lambda item: str(item.get("verification_state") or ""),
    }.get(sort_by, lambda item: str(item.get("created_at") or ""))
    documents.sort(key=sort_key, reverse=sort_dir.lower() == "desc")
    if not paginated:
        return documents
    safe_page = max(page, 1)
    safe_size = min(max(page_size, 1), 100)
    return list_page(documents, safe_page, safe_size)


@router.get("/notifications")
def list_notifications(unread_only: bool = False, user: User = Depends(current_user), db: Session = Depends(get_db)):
    query = select(InAppNotification).where(InAppNotification.recipient_id == user.id)
    if unread_only: query = query.where(InAppNotification.read_at.is_(None))
    return [model_dict(item, [field.name for field in item.__table__.columns]) for item in db.scalars(query.order_by(InAppNotification.created_at.desc()))]


@router.post("/notifications/{notification_id}/read")
def read_notification(notification_id: str, user: User = Depends(current_user), db: Session = Depends(get_db)):
    notification = db.get(InAppNotification, notification_id)
    if not notification or notification.recipient_id != user.id: raise HTTPException(404, "Notification not found")
    notification.read_at = datetime.now(UTC)
    db.commit(); return {"ok": True}


@router.delete("/documents/{evidence_id}")
def delete_document(evidence_id: str, user: User = Depends(current_user), db: Session = Depends(get_db)):
    evidence = db.get(ComplianceEvidence, evidence_id)
    if not evidence:
        raise HTTPException(404, "Document not found")
    instance = get_instance_or_404(db, user, evidence.instance_id)
    require_manager(db, user, instance.organization_id)
    stored_filename = evidence.stored_filename
    audit(db, organization_id=instance.organization_id, actor_id=user.id, action="DELETE_DOCUMENT", module="EVIDENCE", entity_type="ComplianceEvidence", entity_id=evidence.id, old={"filename": evidence.original_filename})
    db.delete(evidence)
    db.commit()
    delete_stored_file(stored_filename)
    return {"ok": True}


@router.get("/reports/compliance-status.csv")
def compliance_status_report(status: str | None = None, user: User = Depends(current_user), db: Session = Depends(get_db)):
    org = org_for(user)
    rows = []
    for instance in db.scalars(select(ComplianceInstance).where(ComplianceInstance.organization_id == org).order_by(ComplianceInstance.due_date)):
        if not can_access_subject(db, user, org, instance.subject_type, instance.subject_id):
            continue
        item = instance_dict(db, instance)
        if status and item["display_status"] != status:
            continue
        rows.append(item)
    output = io.StringIO()
    writer = csv.DictWriter(output, fieldnames=["compliance_id", "compliance_name", "subject_name", "subject_type", "frequency", "due_date", "risk_level", "display_status"])
    writer.writeheader()
    writer.writerows([{key: row.get(key) for key in writer.fieldnames} for row in rows])
    return StreamingResponse(iter([output.getvalue()]), media_type="text/csv", headers={"Content-Disposition": "attachment; filename=compliance-status-report.csv"})


@router.get("/reports/overdue-compliance.csv")
def overdue_compliance_report(user: User = Depends(current_user), db: Session = Depends(get_db)):
    return compliance_status_report("OVERDUE", user, db)


@router.get("/reports/completion.csv")
def completion_report(user: User = Depends(current_user), db: Session = Depends(get_db)):
    return compliance_status_report("COMPLETED", user, db)


@router.get("/reports/entity-compliance.csv")
def entity_compliance_report(user: User = Depends(current_user), db: Session = Depends(get_db)):
    return compliance_status_report(None, user, db)


@router.get("/dashboard/status")
def dashboard_status(user: User = Depends(current_user), db: Session = Depends(get_db)):
    org = org_for(user)
    counts: dict[str, int] = {}
    for instance in db.scalars(select(ComplianceInstance).where(ComplianceInstance.organization_id == org)):
        if can_access_subject(db, user, org, instance.subject_type, instance.subject_id):
            label = instance_dict(db, instance)["display_status"]
            counts[label] = counts.get(label, 0) + 1
    return counts


def dashboard_display_status(instance: ComplianceInstance, today_value: date) -> str:
    if instance.status in {"PENDING", "IN_PROGRESS", "CORRECTION_REQUIRED"}:
        return "OVERDUE" if instance.due_date < today_value else "DUE"
    if instance.status in {"SUBMITTED", "UNDER_REVIEW"}:
        return "PENDING_FOR_APPROVAL"
    if instance.status == "APPROVED":
        return "COMPLETED_LATE" if instance.completion_late else "COMPLETED"
    if instance.status == "NOT_APPLICABLE":
        return "NOT_APPLICABLE"
    return "REJECTED_BY_CHECKER"


@router.get("/dashboard/overview")
def dashboard_overview(
    unit_id: str | None = None,
    entity_id: str | None = None,
    contractor_id: str | None = None,
    period_key: str | None = None,
    user: User = Depends(current_user),
    db: Session = Depends(get_db),
):
    """Return one tenant-scoped, filter-aware dashboard projection.

    The API resolves hierarchy server-side instead of trusting client supplied
    identifiers. Entity is MORAX's existing compliance subject: Unit,
    Contractor, or Contractor Site.
    """
    organization_id = org_for(user)
    today_value = date.today()
    units = {
        item.id: item
        for item in db.scalars(
            select(Unit).where(Unit.organization_id == organization_id)
        )
        if can_access_subject(db, user, organization_id, "UNIT", item.id)
    }
    contractors = {
        item.id: item
        for item in db.scalars(
            select(Contractor).where(Contractor.organization_id == organization_id)
        )
        if can_access_subject(db, user, organization_id, "CONTRACTOR", item.id)
    }
    sites = {
        item.id: item
        for item in db.scalars(
            select(ContractorSite).where(
                ContractorSite.organization_id == organization_id
            )
        )
        if can_access_subject(db, user, organization_id, "CONTRACTOR_SITE", item.id)
    }

    def site_unit_id(site: ContractorSite) -> str | None:
        return site.unit_id or (
            contractors[site.contractor_id].unit_id
            if site.contractor_id in contractors
            else None
        )

    entity_options = [
        {
            "id": item.id,
            "name": item.name,
            "entity_type": "UNIT",
            "unit_id": item.id,
            "contractor_id": None,
        }
        for item in units.values()
    ] + [
        {
            "id": item.id,
            "name": item.name,
            "entity_type": "CONTRACTOR",
            "unit_id": item.unit_id,
            "contractor_id": item.id,
        }
        for item in contractors.values()
    ] + [
        {
            "id": item.id,
            "name": item.name,
            "entity_type": "CONTRACTOR_SITE",
            "unit_id": site_unit_id(item),
            "contractor_id": item.contractor_id,
        }
        for item in sites.values()
    ]

    def matches_context(
        subject_type: str,
        subject_id: str,
        resolved_unit_id: str | None,
        resolved_contractor_id: str | None,
    ) -> bool:
        if unit_id and resolved_unit_id != unit_id:
            return False
        if contractor_id and resolved_contractor_id != contractor_id:
            return False
        if entity_id and subject_id != entity_id:
            return False
        return True

    context_entities = [
        item
        for item in entity_options
        if matches_context(
            item["entity_type"],
            item["id"],
            item["unit_id"],
            item["contractor_id"],
        )
    ]
    context_contractors = [
        item
        for item in contractors.values()
        if (not unit_id or item.unit_id == unit_id)
        and (
            not entity_id
            or any(
                entity["id"] == entity_id
                and (
                    entity["entity_type"] == "UNIT"
                    or entity["contractor_id"] == item.id
                )
                for entity in entity_options
            )
        )
    ]

    rows: list[dict] = []
    for instance in db.scalars(
        select(ComplianceInstance)
        .where(ComplianceInstance.organization_id == organization_id)
        .order_by(ComplianceInstance.due_date)
    ):
        if not can_access_subject(
            db, user, organization_id, instance.subject_type, instance.subject_id
        ):
            continue
        subject = (
            units.get(instance.subject_id)
            if instance.subject_type == "UNIT"
            else contractors.get(instance.subject_id)
            if instance.subject_type == "CONTRACTOR"
            else sites.get(instance.subject_id)
        )
        if not subject:
            continue
        resolved_unit_id = (
            instance.subject_id
            if instance.subject_type == "UNIT"
            else subject.unit_id
            if instance.subject_type == "CONTRACTOR"
            else site_unit_id(subject)
        )
        resolved_contractor_id = (
            instance.subject_id
            if instance.subject_type == "CONTRACTOR"
            else subject.contractor_id
            if instance.subject_type == "CONTRACTOR_SITE"
            else None
        )
        if not matches_context(
            instance.subject_type,
            instance.subject_id,
            resolved_unit_id,
            resolved_contractor_id,
        ):
            continue
        if period_key and instance.period_key != period_key:
            continue
        snapshot = json.loads(instance.rule_snapshot_json)
        display_status = dashboard_display_status(instance, today_value)
        rows.append(
            {
                "id": instance.id,
                "compliance_id": snapshot.get("compliance_id"),
                "compliance_name": snapshot.get("name"),
                "frequency": snapshot.get("frequency"),
                "risk_level": snapshot.get("risk_level"),
                "subject_name": subject.name,
                "subject_type": instance.subject_type,
                "subject_id": instance.subject_id,
                "unit_id": resolved_unit_id,
                "contractor_id": resolved_contractor_id,
                "period_key": instance.period_key,
                "due_date": instance.due_date.isoformat(),
                "status": instance.status,
                "display_status": display_status,
                "row_version": instance.row_version,
                "is_overdue": display_status == "OVERDUE",
                "days_overdue": max(0, (today_value - instance.due_date).days)
                if display_status == "OVERDUE"
                else 0,
                "required_document": snapshot.get("required_document"),
            }
        )

    total = len(rows)
    completed = sum(item["status"] == "APPROVED" for item in rows)
    overdue = sum(item["display_status"] == "OVERDUE" for item in rows)
    due_soon = sum(
        item["status"] not in {"APPROVED", "NOT_APPLICABLE"}
        and today_value <= date.fromisoformat(item["due_date"]) <= today_value + timedelta(days=7)
        for item in rows
    )
    pending = sum(
        item["status"] in {"PENDING", "IN_PROGRESS", "CORRECTION_REQUIRED"}
        for item in rows
    )
    statuses: dict[str, int] = {}
    frequencies: dict[str, int] = {}
    entity_health: dict[str, dict] = {}
    contractor_health: dict[str, dict] = {}
    for item in rows:
        status = item["display_status"]
        statuses[status] = statuses.get(status, 0) + 1
        frequency = item["frequency"] or "UNCONFIGURED"
        frequencies[frequency] = frequencies.get(frequency, 0) + 1
        entity = entity_health.setdefault(
            item["subject_id"],
            {
                "id": item["subject_id"],
                "name": item["subject_name"],
                "entity_type": item["subject_type"],
                "total": 0,
                "completed": 0,
                "overdue": 0,
            },
        )
        entity["total"] += 1
        entity["completed"] += item["status"] == "APPROVED"
        entity["overdue"] += item["display_status"] == "OVERDUE"
        if item["contractor_id"]:
            contractor = contractors.get(item["contractor_id"])
            if contractor:
                contractor_data = contractor_health.setdefault(
                    contractor.id,
                    {
                        "id": contractor.id,
                        "name": contractor.name,
                        "total": 0,
                        "completed": 0,
                        "overdue": 0,
                    },
                )
                contractor_data["total"] += 1
                contractor_data["completed"] += item["status"] == "APPROVED"
                contractor_data["overdue"] += item["display_status"] == "OVERDUE"

    def health_rows(values: dict[str, dict]) -> list[dict]:
        output = []
        for item in values.values():
            rate = round(item["completed"] / item["total"] * 100, 1) if item["total"] else 0
            output.append({**item, "rate": rate})
        return sorted(output, key=lambda item: (item["rate"], -item["overdue"], item["name"]))

    periods = sorted({item["period_key"] for item in rows}, reverse=True)
    upcoming = [
        item
        for item in rows
        if item["status"] not in {"APPROVED", "NOT_APPLICABLE"}
        and today_value <= date.fromisoformat(item["due_date"]) <= today_value + timedelta(days=14)
    ][:8]
    return {
        "filters": {
            "units": [
                {"id": item.id, "name": item.name}
                for item in sorted(units.values(), key=lambda row: row.name)
            ],
            "entities": sorted(context_entities, key=lambda row: (row["name"], row["entity_type"])),
            "contractors": [
                {"id": item.id, "name": item.name, "unit_id": item.unit_id}
                for item in sorted(context_contractors, key=lambda row: row.name)
            ],
            "periods": periods,
        },
        "summary": {
            "total": total,
            "completed": completed,
            "pending": pending,
            "overdue": overdue,
            "due_soon": due_soon,
            "completion_rate": round(completed / total * 100, 1) if total else 0,
        },
        "status_distribution": [
            {"name": name, "value": value}
            for name, value in sorted(statuses.items())
        ],
        "frequency_distribution": [
            {"name": name, "value": value}
            for name, value in sorted(frequencies.items())
        ],
        "entity_health": health_rows(entity_health),
        "contractor_health": health_rows(contractor_health),
        "upcoming": upcoming,
    }


@router.get("/dashboard/frequency")
def dashboard_frequency(user: User = Depends(current_user), db: Session = Depends(get_db)):
    org = org_for(user)
    counts: dict[str, int] = {}
    for instance in db.scalars(select(ComplianceInstance).where(ComplianceInstance.organization_id == org)):
        if can_access_subject(db, user, org, instance.subject_type, instance.subject_id):
            frequency = instance_dict(db, instance)["frequency"]
            counts[frequency] = counts.get(frequency, 0) + 1
    return counts


@router.get("/dashboard/entities")
def dashboard_entities(user: User = Depends(current_user), db: Session = Depends(get_db)):
    org = org_for(user)
    summary: dict[str, dict[str, int]] = {}
    for instance in db.scalars(select(ComplianceInstance).where(ComplianceInstance.organization_id == org)):
        if not can_access_subject(db, user, org, instance.subject_type, instance.subject_id):
            continue
        row = instance_dict(db, instance)
        entity = summary.setdefault(row["subject_name"], {"due": 0, "overdue": 0, "done": 0, "total": 0})
        entity["total"] += 1
        if row["display_status"] == "OVERDUE": entity["overdue"] += 1
        elif row["display_status"] in {"DUE", "PENDING_FOR_APPROVAL"}: entity["due"] += 1
        else: entity["done"] += 1
    return [{"entity": name, **values, "rate": round(values["done"] / values["total"] * 100, 1) if values["total"] else 0} for name, values in summary.items()]


@router.get("/dashboard/summary")
def dashboard(user: User = Depends(current_user), db: Session = Depends(get_db)):
    org = org_for(user); rows = [item for item in db.scalars(select(ComplianceInstance).where(ComplianceInstance.organization_id == org)) if can_access_subject(db, user, org, item.subject_type, item.subject_id)]
    statuses = {key: 0 for key in ["PENDING", "IN_PROGRESS", "SUBMITTED", "UNDER_REVIEW", "CORRECTION_REQUIRED", "REJECTED", "APPROVED"]}
    for row in rows: statuses[row.status] = statuses.get(row.status, 0) + 1
    overdue = sum(1 for row in rows if row.status != "APPROVED" and row.due_date < date.today())
    completed_late = sum(1 for row in rows if row.status == "APPROVED" and row.completion_late)
    return {"total": len(rows), "statuses": statuses, "overdue": overdue, "completed_late": completed_late}


@router.get("/audit-logs")
def audit_logs(page: int = 1, page_size: int = 100, user: User = Depends(current_user), db: Session = Depends(get_db)):
    org = org_for(user)
    if not can_read_audit(db, user, org):
        raise HTTPException(403, "Organization administrator or organization-scoped auditor permission is required")
    rows = [model_dict(item, [field.name for field in item.__table__.columns]) for item in db.scalars(select(AuditLog).where(AuditLog.organization_id == org).order_by(AuditLog.created_at.desc()))]
    return list_page(rows, page, min(page_size, 100))
